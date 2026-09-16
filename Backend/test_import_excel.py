import io
import tempfile

from fastapi.testclient import TestClient

import main as app_main
from database import Base, engine, SessionLocal
import models

client = TestClient(app_main.app)


def test_health():
    r = client.get("/health")
    assert r.status_code == 200, r.text
    assert r.json() == {"estado": "ok"}


def test_plantilla():
    r = client.get("/api/v1/productos/plantilla")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats"
    )
    assert r.content[:2] == b"PK"


def test_importar_ok():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre", "Precio", "Stock", "Categoria", "Descripcion"])
    ws.append(["Producto A", 10.5, 3, "Cat1", "desc A"])
    ws.append(["Producto B", 20, 5, "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    r = client.post(
        "/api/v1/productos/importar",
        files={"file": ("inv.xlsx", buf, "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["importados"] == 2, data
    assert data["errores"] == [], data
    assert data["total_filas"] == 2, data

    # verificar en BD
    db = SessionLocal()
    try:
        a = db.query(models.Producto).filter(models.Producto.nombre == "Producto A").first()
        b = db.query(models.Producto).filter(models.Producto.nombre == "Producto B").first()
        assert a is not None and a.precio == 10.5 and a.stock == 3
        assert b is not None and b.stock == 5
    finally:
        db.close()


def test_importar_upsert():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre", "Precio", "Stock", "Categoria", "Descripcion"])
    ws.append(["Producto A", 99, 7, "CatX", "actualizado"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = client.post(
        "/api/v1/productos/importar",
        files={"file": ("inv2.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["importados"] == 1, data
    db = SessionLocal()
    try:
        a = db.query(models.Producto).filter(models.Producto.nombre == "Producto A").first()
        assert a is not None and a.precio == 99 and a.stock == 7 and a.categoria == "CatX"
    finally:
        db.close()


def test_importar_sin_columnas():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Otro"])
    ws.append(["x"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = client.post(
        "/api/v1/productos/importar",
        files={"file": ("inv.xlsx", buf, "application/octet-stream")},
    )
    assert r.status_code == 400, r.text


def test_importar_no_excel():
    r = client.post(
        "/api/v1/productos/importar",
        files={"file": ("datos.csv", b"a,b", "text/csv")},
    )
    assert r.status_code == 400, r.text


def test_importar_acierta_errores_por_fila():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Nombre", "Precio", "Stock", "Categoria", "Descripcion"])
    ws.append(["", 5, 1, "c", ""])          # nombre vacio -> error
    ws.append(["Valido", -1, 1, "c", ""])   # precio negativo -> error
    ws.append(["Bueno", 2, 3, "c", "ok"])   # ok
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    r = client.post(
        "/api/v1/productos/importar",
        files={"file": ("inv.xlsx", buf, "application/octet-stream")},
    )
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["importados"] == 1, data
    assert len(data["errores"]) == 2, data
    assert data["total_filas"] == 3, data


if __name__ == "__main__":
    import traceback
    def run(name, fn):
        try:
            fn()
            print(f"PASS {name}")
        except Exception:
            print(f"FAIL {name}")
            traceback.print_exc()

    run("health", test_health)
    run("plantilla", test_plantilla)
    run("importar_ok", test_importar_ok)
    run("importar_upsert", test_importar_upsert)
    run("importar_sin_columnas", test_importar_sin_columnas)
    run("importar_no_excel", test_importar_no_excel)
    run("importar_errores_por_fila", test_importar_acierta_errores_por_fila)