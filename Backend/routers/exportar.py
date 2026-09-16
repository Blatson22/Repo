"""Exportación de reportes a CSV, XLSX y PDF.

Endpoints:
    GET /api/v1/exportar/reportes?reporte=movimientos&formato=csv
    GET /api/v1/exportar/reportes?reporte=existencias&formato=xlsx
    GET /api/v1/exportar/reportes?reporte=top&tipo=venta&formato=pdf

repor_type: resumen | movimientos | existencias | top | serie
formato:    csv | xlsx | pdf
"""
import csv
import io
from datetime import date, datetime, time
from typing import List, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

import models
from database import get_db

router = APIRouter(
    prefix="/api/v1/exportar",
    tags=["exportar"],
)


def _titulo(reporte: str, tipo: Optional[str] = None) -> str:
    nombres = {
        "resumen": "Resumen general",
        "movimientos": "Movimientos",
        "existencias": "Existencias",
        "top": "Top productos",
        "serie": "Serie por periodos",
    }
    t = nombres.get(reporte, reporte)
    if tipo:
        t = f"{t} ({tipo})"
    return t


def _filtros_fecha(fecha_desde, fecha_hasta, q):
    if fecha_desde:
        q = q.filter(models.Documento.fecha >= datetime.combine(fecha_desde, time.min))
    if fecha_hasta:
        q = q.filter(models.Documento.fecha <= datetime.combine(fecha_hasta, time.max))
    return q


# En cada tipo de reporte devolvemos (cabeceras, filas de textos).
def _datos(reporte: str, db: Session, tipo, fecha_desde, fecha_hasta, producto_id, agregar_por) -> tuple[List[str], List[List[str]]]:
    cabeceras: List[str] = []
    filas: List[List[str]] = []

    if reporte == "resumen":
        q = db.query(models.Documento.tipo, func.coalesce(func.sum(models.Documento.total), 0.0))
        q = _filtros_fecha(fecha_desde, fecha_hasta, q).group_by(models.Documento.tipo)
        por_tipo = {t: v for t, v in q.all()}
        valor = db.query(func.coalesce(func.sum(models.Producto.precio * models.Producto.stock), 0.0)).scalar()
        conteos = {
            t: db.query(func.count(models.Documento.id)).filter(models.Documento.tipo == t).scalar() or 0
            for t in models.TIPOS_DOCUMENTO
        }
        cabeceras = ["Concepto", "Importe", "Cantidad documentos"]
        filas = [
            ["Compras", round(por_tipo.get("compra", 0), 2), conteos["compra"]],
            ["Ventas", round(por_tipo.get("venta", 0), 2), conteos["venta"]],
            ["Despachos", round(por_tipo.get("despacho", 0), 2), conteos["despacho"]],
            ["Ajustes", round(por_tipo.get("ajuste", 0), 2), conteos["ajuste"]],
            ["Valor del almacen", round(valor or 0, 2), ""],
        ]
        return cabeceras, [[str(c) for c in f] for f in filas]

    if reporte == "movimientos":
        q = (
            db.query(
                models.Documento.id,
                models.Documento.folio,
                models.Documento.tipo,
                models.Documento.fecha,
                models.Producto.nombre,
                models.DocumentoLinea.cantidad,
                models.DocumentoLinea.signo,
                models.DocumentoLinea.precio_unitario,
                models.DocumentoLinea.subtotal,
            )
            .join(models.Documento, models.Documento.id == models.DocumentoLinea.documento_id)
            .join(models.Producto, models.Producto.id == models.DocumentoLinea.producto_id)
        )
        if producto_id:
            q = q.filter(models.DocumentoLinea.producto_id == producto_id)
        if tipo:
            q = q.filter(models.Documento.tipo == tipo)
        q = _filtros_fecha(fecha_desde, fecha_hasta, q)
        rows = q.order_by(models.Documento.fecha.desc()).all()
        cabeceras = ["Documento", "Folio", "Tipo", "Fecha", "Producto", "Cantidad", "Signo", "Precio", "Subtotal"]
        filas = [
            [
                str(r[0]), r[1], r[2], r[3].strftime("%Y-%m-%d %H:%M"),
                r[4], str(r[5]), str(r[6]), str(r[7]), str(r[8]),
            ]
            for r in rows
        ]
        return cabeceras, filas

    if reporte == "existencias":
        q = db.query(models.Producto).order_by(models.Producto.nombre)
        rows = q.all()
        cabeceras = ["ID", "Nombre", "Categoria", "Stock", "Precio", "Valor"]
        filas = [
            [str(p.id), p.nombre, p.categoria or "", str(p.stock or 0), str(p.precio or 0),
             str(round((p.precio or 0) * (p.stock or 0), 2))]
            for p in rows
        ]
        return cabeceras, filas

    if reporte == "top":
        tipo_req = tipo or "venta"
        q = (
            db.query(
                models.DocumentoLinea.producto_id,
                models.Producto.nombre,
                func.sum(models.DocumentoLinea.cantidad).label("cant"),
                func.sum(models.DocumentoLinea.subtotal).label("imp"),
            )
            .join(models.Producto, models.Producto.id == models.DocumentoLinea.producto_id)
            .join(models.Documento, models.Documento.id == models.DocumentoLinea.documento_id)
            .filter(models.Documento.tipo == tipo_req)
        )
        q = _filtros_fecha(fecha_desde, fecha_hasta, q)
        rows = (
            q.group_by(models.DocumentoLinea.producto_id, models.Producto.nombre)
            .order_by(func.sum(models.DocumentoLinea.cantidad).desc())
            .all()
        )
        cabeceras = ["Producto", "Cantidad", "Importe"]
        filas = [[r.nombre, str(r.cant or 0), str(round(r.imp or 0, 2))] for r in rows]
        return cabeceras, filas

    if reporte == "serie":
        tipo_req = tipo or "venta"
        q = db.query(models.Documento.fecha, models.Documento.total).filter(models.Documento.tipo == tipo_req)
        q = _filtros_fecha(fecha_desde, fecha_hasta, q)
        docs = q.all()
        grupo: dict[str, list] = {}
        for fecha_val, total in docs:
            clave = fecha_val.strftime("%Y-%m") if agregar_por == "mes" else fecha_val.strftime("%Y-%m-%d")
            g = grupo.setdefault(clave, [0.0, 0])
            g[0] += total or 0.0
            g[1] += 1
        cabeceras = ["Periodo", "Total", "Documentos"]
        filas = [[k, str(round(v[0], 2)), str(v[1])] for k, v in sorted(grupo.items())]
        return cabeceras, filas

    raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Reporte no soportado")


def _csv(cabeceras, filas, titulo) -> StreamingResponse:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([titulo])
    writer.writerow([])
    writer.writerow(cabeceras)
    writer.writerows(filas)
    data = buffer.getvalue().encode("utf-8-sig")
    return StreamingResponse(
        iter([data]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": "attachment; filename=reporte.csv"},
    )


def _xlsx(cabeceras, filas, titulo) -> StreamingResponse:
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws.append([titulo])
    ws.append([])
    ws.append(cabeceras)
    for fila in filas:
        ws.append(fila)
    for i in range(1, len(cabeceras) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 22
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=reporte.xlsx"},
    )


def _pdf(cabeceras, filas, titulo) -> StreamingResponse:
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=10 * mm, rightMargin=10 * mm)
    story = []
    estilos = getSampleStyleSheet()
    story.append(Paragraph(f"<b>{titulo}</b>", estilos["Title"]))
    story.append(Spacer(1, 6 * mm))
    tabla = Table([cabeceras] + filas, repeatRows=1)
    tabla.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4b6584")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f2f6")]),
            ]
        )
    )
    story.append(tabla)
    doc.build(story)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=reporte.pdf"},
    )


@router.get("/reportes")
def exportar_reporte(
    reporte: str = Query(..., description="resumen|movimientos|existencias|top|serie"),
    formato: str = Query("csv", pattern="^(csv|xlsx|pdf)$"),
    tipo: Optional[str] = Query(None),
    producto_id: Optional[int] = Query(None),
    agregar_por: str = Query("dia", pattern="^(dia|mes)$"),
    fecha_desde: Optional[date] = Query(None),
    fecha_hasta: Optional[date] = Query(None),
    db: Session = Depends(get_db),
):
    if reporte in ("top", "serie") and tipo not in ("venta", "despacho", "compra"):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Para top/serie se requiere tipo=venta|despacho|compra",
        )
    cabeceras, filas = _datos(
        reporte, db, tipo, fecha_desde, fecha_hasta, producto_id, agregar_por
    )
    titulo = _titulo(reporte, tipo)
    if formato == "csv":
        return _csv(cabeceras, filas, titulo)
    if formato == "xlsx":
        return _xlsx(cabeceras, filas, titulo)
    return _pdf(cabeceras, filas, titulo)